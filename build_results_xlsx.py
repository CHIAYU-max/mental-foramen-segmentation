import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 把目前所有實驗結果整理成一個 Excel 檔。
# 數字取自各分析程式的輸出(每張表下方註明來源);衍生的數值(平均、標準差、差距)以公式計算。
# 附錄工作表直接複製結果 CSV 的原始內容。
ROOT = Path(__file__).resolve().parent
OUT = ROOT / "compare_results" / "results_summary.xlsx"
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
D4, D3, PCT, P = "0.0000", "0.000", "0.0%", "0.00E+00"


class Sheet:
    def __init__(self, wb, title, heading, intro=()):
        self.ws = wb.create_sheet(title)
        self.row = 1
        self.widths = {}
        self.cell(1, heading, bold=True, size=14)
        self.row = 2
        for line in intro:
            self.cell(1, line, italic=True)
            self.row += 1
        self.row += 1

    def cell(self, col, value, bold=False, italic=False, size=10, fmt=None, fill=None):
        c = self.ws.cell(row=self.row, column=col, value=value)
        c.font = Font(name=FONT, bold=bold, italic=italic, size=size)
        if fmt and not isinstance(value, str) or (fmt and isinstance(value, str) and value.startswith("=")):
            c.number_format = fmt
        if fill:
            c.fill = fill
        c.alignment = Alignment(vertical="center")
        if isinstance(value, str) and not value.startswith("="):
            self.widths[col] = max(self.widths.get(col, 0), min(len(value) * 1.6 + 2, 60))
        return c

    def table(self, title, header, rows_fn, fmts=None, note=None):
        # rows_fn 接收第一列資料的列號,回傳各列;公式中的 {r} 會換成該列的列號
        self.cell(1, title, bold=True, size=11)
        self.row += 1
        for i, h in enumerate(header, start=1):
            self.cell(i, h, bold=True, fill=HEADER_FILL)
        self.row += 1
        first = self.row
        for values in rows_fn(first):
            for i, v in enumerate(values, start=1):
                if isinstance(v, str) and v.startswith("="):
                    v = v.format(r=self.row)
                fmt = fmts[i - 1] if fmts and i - 1 < len(fmts) else None
                self.cell(i, v, fmt=fmt)
            self.row += 1
        if note:
            self.cell(1, note, italic=True)
            self.row += 1
        self.row += 1
        return first

    def finish(self):
        for col, width in self.widths.items():
            self.ws.column_dimensions[get_column_letter(col)].width = max(10, width)
        self.ws.freeze_panes = None


def overview(wb):
    s = Sheet(wb, "說明", "頦孔與下齒槽神經分割:實驗結果總表")
    rows = [
        ("目的", "在全景片 ROI 上分割頦孔(mf)與下齒槽神經管(ian),比較 YOLOv8 實例分割與 U-Net 語意分割"),
        ("訓練資料", "559 張 ROI(第一批 297 張、第二批 262 張),裁自至少 294 張全景片(另 29 張找不到來源)。"
                     "舊稿的「341 位病人」是原始檔案編號數,不是病人數(count_patients.py,2026-10-26),不能使用"),
        ("獨立測試集", "新切分 137 張 ROI、103 張來源全景片(去掉與訓練同一張全景片的 209_Left_roi);舊切分 138 張"),
        ("新切分(論文主表)", "病人編號 ∪ 來源全景片以 union-find 合併成 301 組,整組放進同一折(kfold_pano.py,2026-10-24);見「新切分_*」工作表"),
        ("ROI 裁切", "依 FDI 33–36 / 43–46,上緣為小臼齒中心、下緣為偵測到的下顎下緣;只裁切不縮放"),
        ("影像前處理", "Unsharp mask(各種前處理在 K 折比較中無差異)"),
        ("切分(舊)", "以病人編號為單位的 5 折交叉驗證(各折 111–112 張),兩批各自分配;除了「新切分_*」以外的工作表都是這個切分"),
        ("最終 U-Net", "ResNet34 encoder(ImageNet 預訓練)+ U-Net decoder;輸入 512×512 灰階 / 255;強增強(翻轉、亮度對比、縮放 0.8–1.2、旋轉 ±15°、平移 ±10%);AdamW lr 3e-4、cosine、100 epochs、batch 8;交叉熵 + Dice loss;取驗證 Dice 最佳的 epoch"),
        ("YOLO", "YOLOv8n-seg,imgsz 1280、200 epochs、patience 50、batch 16;推論時每個類別取信心最高的候選(conf ≥ 0.01、NMS IoU 0.3)"),
        ("多數決", "測試集上 5 折模型逐像素多數決"),
        ("後處理(lcc)", "2026-10-28 起頦孔只保留最大連通區塊(8 連通,多數決之後);YOLO 本來就只取一個實例,不受影響。「新切分_*」中標 lcc 的表為論文主結果"),
        ("指標", "Dice、IoU、容許 2/4 像素邊界誤差的 Dice、HD95、ASSD;頦孔定位:命中率 = 預測中心落在 GT 等效半徑 R 以內的比例"),
        ("信賴區間", "以病人為單位重抽樣 2000 次的 95% 區間(種子 0)"),
        ("注意", "沒有像素間距,距離以像素或 R 的倍數表示;兩批全景片解析度不一致,以 R 為單位的指標較可比"),
        ("尚未完成", "同一病人不同時期影像的人工確認(find_same_patient.py)"),
    ]
    s.table("研究設定", ["項目", "內容"], lambda first: rows)
    sheets = [
        ("新切分_主結果", "新切分: 七個模型、U-Net 與 YOLO26s 五個種子、合併種子的病人層級檢定(論文主表;以 lcc 版本為準)"),
        ("新切分_延伸", "新切分(lcc): 困難案例、可辨識性分層、串接式推論、R34/384、推論時翻轉、半監督"),
        ("新切分_運算量", "新切分: U-Net encoder 與輸入尺寸的精度對運算量(三個種子)"),
        ("主要結果_種子","最終 U-Net 與 YOLO 在三個隨機種子下的結果,平均、標準差與差距以公式計算"),
        ("差值CI", "U-Net − YOLO 差值的 95% 信賴區間(各種子、三種子平均、依可辨識性分組)"),
        ("詳細指標", "種子 0 的完整指標與 95% 信賴區間"),
        ("消融實驗", "結構、預訓練、增強各自的貢獻"),
        ("學習曲線", "訓練資料量 25%–100% 的表現"),
        ("批次分析", "兩批資料的差距、交叉批次實驗、標註特徵、頦孔暗度分組"),
        ("標註一致性", "40 張重新標註的結果"),
        ("其他實驗", "推論規則、候選挑選、模型大小、設計變體、切分方式、前處理"),
        ("可辨識性", "頦孔可辨識的比例,以及依可辨識性分層的結果"),
        ("前處理", "9 種前處理與原圖的比較(強增強與無增強)"),
        ("可信度", "各實驗的可信度分級(A/B/C)與限制"),
        ("附錄_*", "結果 CSV 的原始內容"),
    ]
    s.table("工作表", ["工作表", "內容"], lambda first: sheets)
    s.finish()


def seeds(wb):
    s = Sheet(wb, "主要結果_種子", "最終 U-Net 與 YOLO:三個隨機種子", ["種子 0 為 kfold_patient.py,種子 1、2 為 seed_replicates.py;切分與設定完全相同"])
    data = {
        "交叉驗證(559 張)": {
            "U-Net": [("mf Dice", [0.7005, 0.7088, 0.7038], D4), ("mf 命中率 1R", [0.8640, 0.8587, 0.8605], PCT), ("ian Dice", [0.6427, 0.6422, 0.6430], D4)],
            "YOLO": [("mf Dice", [0.6783, 0.6690, 0.6686], D4), ("mf 命中率 1R", [0.8193, 0.7907, 0.7907], PCT), ("ian Dice", [0.6111, 0.5937, 0.6033], D4)],
        },
        "測試集(138 張,5 折多數決)": {
            "U-Net": [("mf Dice", [0.6890, 0.7191, 0.7147], D4), ("mf 命中率 1R", [0.9203, 0.9348, 0.9058], PCT), ("ian Dice", [0.6706, 0.6858, 0.6777], D4)],
            "YOLO": [("mf Dice", [0.6725, 0.6636, 0.6141], D4), ("mf 命中率 1R", [0.8551, 0.8261, 0.7681], PCT), ("ian Dice", [0.6335, 0.6443, 0.6372], D4)],
        },
    }
    header = ["模型", "指標", "seed 0", "seed 1", "seed 2", "平均", "標準差"]
    for title, models in data.items():
        def rows(first, models=models):
            out = []
            for name in ("U-Net", "YOLO"):
                for metric, values, _ in models[name]:
                    out.append([name, metric, *values, "=AVERAGE(C{r}:E{r})", "=STDEV(C{r}:E{r})"])
            for i, (metric, _, _) in enumerate(models["U-Net"]):
                u, y = first + i, first + 3 + i
                out.append(["U-Net − YOLO", metric] + [f"={c}{u}-{c}{y}" for c in "CDE"] + ["=AVERAGE(C{r}:E{r})", "=STDEV(C{r}:E{r})"])
            return out
        s.table(title, header, rows, [None, None, D4, D4, D4, D4, D4], note="命中率以小數表示(0.86 = 86%)。來源: seed_replicates.py --eval-only 的輸出")
    s.finish()


def details(wb):
    s = Sheet(wb, "詳細指標", "種子 0 的完整指標(95% 信賴區間以病人為單位重抽樣)")
    header = ["設定", "模型", "mf Dice", "mf Dice 95% CI", "mf IoU", "Dice ±2px", "Dice ±4px", "mf 有預測", "命中 0.5R", "命中 1R", "命中 1R 95% CI",
              "中心誤差/R(中位數)", "HD95/R(中位數)", "中心誤差 px", "HD95 px", "ASSD px", "ian Dice", "ian Dice 95% CI", "ian IoU", "ian Dice ±4px", "ian HD95 px"]
    rows = [
        ["交叉驗證", "YOLOv8n-seg", 0.678, "0.649–0.706", 0.584, 0.744, 0.776, 0.970, 0.696, 0.819, "0.783–0.853", 0.197, 0.420, 3.216, 6.846, 2.213, 0.611, "0.591–0.630", 0.471, 0.761, 23.397],
        ["交叉驗證", "U-Net ResNet34", 0.701, "0.677–0.724", 0.593, 0.774, 0.812, 0.979, 0.692, 0.864, "0.835–0.892", 0.258, 0.530, 4.069, 8.336, 2.619, 0.643, "0.623–0.662", 0.505, 0.782, 19.802],
        ["測試集 單一模型平均", "YOLOv8n-seg", 0.640, "0.592–0.686", 0.539, 0.705, 0.741, 0.914, 0.645, 0.800, "0.741–0.853", 0.323, 0.632, 6.192, 12.348, 4.117, 0.608, "0.576–0.639", 0.462, 0.754, 32.053],
        ["測試集 單一模型平均", "U-Net ResNet34", 0.670, "0.624–0.711", 0.555, 0.742, 0.787, 0.954, 0.683, 0.874, "0.829–0.915", 0.343, 0.673, 6.184, 12.941, 4.154, 0.652, "0.623–0.678", 0.506, 0.798, 25.300],
        ["測試集 5 折多數決", "YOLOv8n-seg", 0.672, "0.617–0.724", 0.568, 0.740, 0.777, 0.920, 0.688, 0.855, "0.783–0.917", 0.252, 0.485, 4.615, 9.009, 2.972, 0.633, "0.599–0.665", 0.488, 0.782, 24.821],
        ["測試集 5 折多數決", "U-Net ResNet34", 0.689, "0.642–0.734", 0.575, 0.765, 0.811, 0.971, 0.717, 0.920, "0.871–0.962", 0.254, 0.525, 4.561, 10.369, 3.400, 0.671, "0.640–0.700", 0.526, 0.818, 16.966],
    ]
    fmts = [None, None, D3, None, D3, D3, D3, PCT, PCT, PCT, None, D3, D3, D3, D3, D3, D3, None, D3, D3, D3]
    s.table("頦孔與下齒槽神經", header, lambda first: rows, fmts, note="距離為像素;兩批解析度不同,像素距離只作參考。來源: paper_metrics.py")
    s.finish()


def ablation(wb):
    s = Sheet(wb, "消融實驗", "消融實驗(病人層級切分、種子 0)", ["差距 = 該版本 − 最終模型,以公式計算;p 為逐張成對 Wilcoxon 檢定(交叉驗證)"])
    header = ["版本", "結構", "預訓練", "增強", "CV mf Dice", "CV mf Dice 95% CI", "CV 命中 1R", "測試 單一 mf Dice", "測試 單一 命中 1R",
              "測試 多數決 mf Dice", "測試 多數決 命中 1R", "CV ian Dice", "測試 多數決 ian Dice", "CV mf 差距", "測試 多數決 mf 差距", "CV ian 差距", "CV mf p", "CV ian p"]
    data = [
        ["最終模型", "ResNet34", "ImageNet", "強", 0.701, "0.677–0.724", 0.864, 0.670, 0.874, 0.689, 0.920, 0.643, 0.671, None, None],
        ["ResNet34 從零", "ResNet34", "無", "強", 0.694, "0.669–0.718", 0.848, 0.675, 0.883, 0.690, 0.913, 0.622, 0.657, 1.02e-01, 5.79e-10],
        ["ResNet34 基本增強", "ResNet34", "ImageNet", "基本", 0.675, "0.649–0.699", 0.832, 0.661, 0.881, 0.680, 0.920, 0.599, 0.653, 8.64e-05, 2.10e-16],
        ["原 U-Net 強增強", "原 U-Net", "無", "強", 0.628, "0.599–0.657", 0.785, 0.543, 0.778, 0.542, 0.797, 0.562, 0.590, 4.80e-18, 1.54e-42],
        ["原 U-Net 基本增強", "原 U-Net", "無", "基本", 0.551, "0.522–0.580", 0.648, 0.474, 0.693, 0.469, 0.775, 0.509, 0.557, 4.01e-43, 2.27e-62],
    ]

    def rows(first):
        out = []
        for values in data:
            base, (p_mf, p_ian) = values[:13], values[13:]
            out.append(base + [f"=E{{r}}-E{first}", f"=J{{r}}-J{first}", f"=L{{r}}-L{first}", p_mf if p_mf is not None else "—", p_ian if p_ian is not None else "—"])
        return out

    fmts = [None] * 4 + [D3, None, PCT, D3, PCT, D3, PCT, D3, D3, D3, D3, D3, P, P]
    s.table("頦孔與下齒槽神經", header, rows, fmts, note="來源: ablation_patient.py 與 paper_metrics.py")
    s.finish()


def learning(wb):
    s = Sheet(wb, "學習曲線", "學習曲線(最終 U-Net 設定,病人層級切分)", ["各資料量的更新步數相同;數值為 5 折平均,SD 為折間標準差"])
    data = [
        [25, 113, 0.6505, 0.0241, 0.5934, 0.0060, 0.5905, 0.0397, 0.6129, 0.0096],
        [50, 225, 0.6762, 0.0205, 0.6230, 0.0066, 0.6380, 0.0210, 0.6291, 0.0142],
        [75, 336, 0.6993, 0.0211, 0.6355, 0.0127, 0.6660, 0.0074, 0.6392, 0.0046],
        [100, 447, 0.7005, 0.0226, 0.6427, 0.0094, 0.6696, 0.0122, 0.6517, 0.0110],
    ]

    def rows(first):
        last = first + len(data) - 1
        return [d + [f"=C{{r}}/C{last}", f"=G{{r}}/G{last}"] for d in data]

    header = ["資料量 %", "訓練張數", "CV mf Dice", "SD", "CV ian Dice", "SD", "測試 mf Dice", "SD", "測試 ian Dice", "SD", "CV mf 相對 100%", "測試 mf 相對 100%"]
    s.table("各資料量的 Dice", header, rows, [None, None] + [D4] * 8 + [PCT, PCT], note="來源: learning_curve.py")
    s.finish()


def batches(wb):
    s = Sheet(wb, "批次分析", "兩批資料的差異(交叉驗證、種子 0)")
    models = [
        ("最終 U-Net", 0.7966, 0.963, 0.5915, 0.752),
        ("YOLOv8n-seg", 0.8086, 0.960, 0.5306, 0.660),
        ("ResNet34 從零", 0.7925, 0.960, 0.5825, 0.721),
        ("ResNet34 基本增強", 0.7876, 0.946, 0.5466, 0.702),
        ("原 U-Net 強增強", 0.7578, 0.919, 0.4799, 0.634),
        ("原 U-Net 基本增強", 0.6846, 0.774, 0.3991, 0.504),
    ]
    s.table("各模型在兩批上的頦孔表現", ["模型", "第一批 Dice", "第一批 命中 1R", "第二批 Dice", "第二批 命中 1R", "Dice 差距"],
            lambda first: [list(m) + ["=B{r}-D{r}"] for m in models], [None, D4, PCT, D4, PCT, D4], note="第一批 297 張、第二批 262 張。來源: paper_metrics.py")

    sub = [
        ("第一批", 297, 0.7966, 0.963, 0.8086, 0.960, 0.6516),
        ("第二批:只在 mf_262 的標註", 138, 0.5658, 0.739, 0.4935, 0.638, 0.6213),
        ("第二批:有 0827_ 檔名的標註", 124, 0.6202, 0.766, 0.5718, 0.685, 0.6454),
    ]
    s.table("第二批依標註來源", ["組別", "n", "U-Net mf Dice", "U-Net 命中 1R", "YOLO mf Dice", "YOLO 命中 1R", "U-Net ian Dice"],
            lambda first: [list(x) for x in sub], [None, None, D4, PCT, D4, PCT, D4], note="解析度、左右側、頦孔大小分組皆無差異。來源: batch_analysis.py")

    cross = [
        ("全部資料", 0.7966, 0.963, 0.5915, 0.752, 0.6890, 0.920),
        ("只用第一批", 0.7937, 0.946, 0.5181, 0.668, 0.6429, 0.862),
        ("只用第二批", 0.7658, 0.919, 0.5528, 0.691, 0.6746, 0.884),
    ]
    s.table("交叉批次實驗(最終 U-Net 設定)", ["訓練資料", "評在第一批 Dice", "第一批 命中 1R", "評在第二批 Dice", "第二批 命中 1R", "測試集 Dice", "測試集 命中 1R"],
            lambda first: [list(x) for x in cross], [None, D4, PCT, D4, PCT, D4, PCT], note="單一批次訓練只用約一半資料(約低 0.02–0.03)。來源: cross_batch.py")

    labels = [
        ("第一批", 297, 812, 3044, 1.9, 0.970, 1.56, 1.13, 10),
        ("第二批", 262, 950, 3198, 2.1, 0.989, 0.94, 0.55, 4),
        ("測試集", 138, 1059, 3210, 1.9, 0.993, 1.25, 0.86, 2),
    ]
    s.table("標註特徵(中位數)", ["組別", "n", "頦孔面積 px", "神經面積 px", "頦孔–神經距離 px", "相連比例", "輪廓暗度", "中心暗度", "被標記張數"],
            lambda first: [list(x) for x in labels], [None, None, "0", "0", "0.0", PCT, "0.00", "0.00", None],
            note="暗度 = (周圍平均 − 內部平均) / 周圍標準差;輪廓暗度受標註畫法影響,中心暗度只用 GT 中心位置。來源: check_labels.py")

    bins = [
        ("−1.27 ~ 0.28", 42, 0.7678, 98, 0.4835, 0.7770, 0.4476),
        ("0.28 ~ 0.80", 66, 0.7828, 73, 0.6389, 0.8024, 0.5962),
        ("0.80 ~ 1.35", 77, 0.7959, 63, 0.6492, 0.8159, 0.5201),
        ("1.35 ~ 3.31", 112, 0.8161, 28, 0.7164, 0.8192, 0.6735),
    ]
    s.table("依中心暗度分組的頦孔 Dice", ["中心暗度", "第一批 n", "U-Net 第一批", "第二批 n", "U-Net 第二批", "YOLO 第一批", "YOLO 第二批", "U-Net 差距", "YOLO 差距"],
            lambda first: [list(b) + ["=C{r}-E{r}", "=F{r}-G{r}"] for b in bins], [None, None, D4, None, D4, D4, D4, D4, D4], note="來源: contrast_analysis.py")

    reg = [
        ("U-Net", 0.2051, 0.0504, 0.1513),
        ("YOLO", 0.2781, 0.1034, 0.2252),
    ]
    s.table("批次差距的迴歸調整(控制暗度與頦孔面積)", ["模型", "原始差距", "控制輪廓暗度後", "控制中心暗度後", "中心暗度解釋的比例"],
            lambda first: [list(x) + ["=1-D{r}/B{r}"] for x in reg], [None, D4, D4, D4, PCT],
            note="輪廓暗度混入了標註畫法的影響,以中心暗度的結果為準")
    s.finish()


def agreement(wb):
    s = Sheet(wb, "標註一致性", "重新標註 40 張(兩批各 20 張,匿名、隨機順序)",
              ["同一位標註者相隔約數月重標;Dice 只計算重新標註時判斷可辨識的影像;模型為最終 U-Net 在交叉驗證中的預測"])
    data = [
        ("mf", "第一批", 18, 2, 0.8284, 0.7920, 0.7823, 0.492, 1.00),
        ("mf", "第二批", 8, 12, 0.7076, 0.7074, 0.6470, 0.742, 1.32),
        ("mf", "合計", 26, 14, 0.7912, 0.7660, 0.7406, 0.353, 1.06),
        ("ian", "第一批", 11, 9, 0.6980, 0.7368, 0.6629, 0.577, 0.91),
        ("ian", "第二批", 9, 11, 0.6797, 0.6987, 0.7298, 0.820, 0.98),
        ("ian", "合計", 20, 20, 0.6898, 0.7197, 0.6930, 0.622, 0.95),
    ]
    header = ["類別", "批次", "n(可辨識)", "無法辨識", "新 vs 舊", "模型 vs 舊", "模型 vs 新", "p(模型 vs 新 − 模型 vs 舊)", "面積比 新/舊", "無法辨識比例", "模型 vs 舊 ÷ 新 vs 舊"]
    s.table("新舊標註與模型的一致性", header, lambda first: [list(d) + ["=D{r}/(C{r}+D{r})", "=F{r}/E{r}"] for d in data],
            [None, None, None, None, D4, D4, D4, "0.000", "0.00", PCT, PCT], note="來源: annotation_agreement.py score")
    s.finish()


def visibility(wb):
    s = Sheet(wb, "可辨識性", "頦孔可辨識性(697 張,匿名、隨機順序判定)",
              ["可辨識 = 能確定是哪一個結構,且能有把握地畫出邊界;與重新標註時對同 40 張的判斷一致率 77.5%、Cohen's kappa 0.56"])
    s.table("可辨識的比例", ["組別", "可辨識", "總數", "比例"], lambda first: [
        ("第一批", 194, 297, "=B{r}/C{r}"), ("第二批", 42, 262, "=B{r}/C{r}"), ("測試集", 58, 138, "=B{r}/C{r}"),
    ], [None, None, None, PCT])
    s.table("交叉驗證:批次 × 可辨識性", ["批次", "組別", "n", "U-Net mf Dice", "U-Net 95% CI", "U-Net 命中 1R", "YOLO mf Dice", "YOLO 95% CI", "YOLO 命中 1R", "U-Net − YOLO"], lambda first: [
        ("第一批", "可辨識", 194, 0.829, "0.809–0.848", 0.969, 0.851, "0.830–0.868", 0.985, "=D{r}-G{r}"),
        ("第一批", "無法辨識", 103, 0.735, "0.703–0.766", 0.951, 0.729, "0.683–0.775", 0.913, "=D{r}-G{r}"),
        ("第二批", "可辨識", 42, 0.822, "0.773–0.864", 0.952, 0.819, "0.777–0.858", 0.976, "=D{r}-G{r}"),
        ("第二批", "無法辨識", 220, 0.548, "0.503–0.592", 0.714, 0.476, "0.426–0.526", 0.600, "=D{r}-G{r}"),
        ("全部", "可辨識", 236, 0.828, "0.809–0.845", 0.966, 0.845, "0.827–0.862", 0.983, "=D{r}-G{r}"),
        ("全部", "無法辨識", 323, 0.607, "0.572–0.643", 0.789, 0.557, "0.516–0.595", 0.700, "=D{r}-G{r}"),
        ("全部", "全部", 559, 0.700, "0.675–0.725", 0.864, 0.678, "0.649–0.707", 0.819, "=D{r}-G{r}"),
    ], [None, None, None, D3, None, PCT, D3, None, PCT, D3], note="種子 0。來源: visibility_grading.py analyze")
    s.table("測試集(5 折多數決)× 可辨識性", ["組別", "n", "U-Net mf Dice", "U-Net 95% CI", "U-Net 命中 1R", "YOLO mf Dice", "YOLO 95% CI", "YOLO 命中 1R", "U-Net − YOLO"], lambda first: [
        ("可辨識", 58, 0.821, "0.793–0.849", 1.000, 0.827, "0.794–0.856", 1.000, "=C{r}-F{r}"),
        ("無法辨識", 80, 0.594, "0.527–0.655", 0.862, 0.561, "0.483–0.640", 0.750, "=C{r}-F{r}"),
        ("全部", 138, 0.689, "0.643–0.730", 0.920, 0.672, "0.618–0.728", 0.855, "=C{r}-F{r}"),
    ], [None, None, D3, None, PCT, D3, None, PCT, D3], note="種子 0。來源: visibility_grading.py analyze")
    s.finish()


def diff_ci(wb):
    # diff_ci.py 的輸出:U-Net − YOLO 差值與 95% 信賴區間
    path = ROOT / "compare_results" / "kfold_patient_mf_559" / "diff_ci.csv"
    if not path.exists():
        print(f"[略過] 找不到 {path},請先執行 python diff_ci.py")
        return
    with open(path, newline="") as f:
        data = list(csv.DictReader(f))
    s = Sheet(wb, "差值CI", "U-Net − YOLO 差值的 95% 信賴區間",
              ["同一張影像上兩模型的指標相減,以病人為單位重抽樣 2000 次;「三種子平均」為每張影像先對三個種子取平均再相減",
               "命中率以比例表示(0.06 = 6 個百分點);CI 不含 0 代表差異顯著"])
    header = ["資料", "種子", "組別", "指標", "n", "U-Net", "YOLO", "差值", "95% CI 下界", "95% CI 上界", "CI 不含 0", "Wilcoxon p"]
    rows = [[r["dataset"], r["seed"], r["group"], r["metric"], int(r["n"]), float(r["unet"]), float(r["yolo"]), float(r["diff"]),
             float(r["ci_low"]), float(r["ci_high"]), '=IF(OR(I{r}>0,J{r}<0),"是","否")', float(r["wilcoxon_p"])] for r in data]
    s.table("各資料集、種子、組別", header, lambda first: rows, [None] * 5 + [D4, D4, D4, D4, D4, None, P], note="來源: diff_ci.py")
    s.finish()


def preprocessing(wb):
    # kfold_preprocessing.py 的輸出
    path = ROOT / "compare_results" / "preprocessing_patient_mf_559" / "summary.csv"
    if not path.exists():
        print(f"[略過] 找不到 {path},請先執行 python kfold_preprocessing.py --eval-only")
        return
    with open(path, newline="") as f:
        data = list(csv.DictReader(f))
    s = Sheet(wb, "前處理", "影像前處理比較(最終 U-Net、病人層級五折、種子 0)",
              ["參數依文獻預設值事先固定;所有版本從同一批原圖出發,只有前處理不同",
               "判定:差異的 95% CI 完全在 ±0.02 內為「實務上等效」;否則 Holm 校正後 p < 0.05 為顯著;其餘為無法判定"])

    def num(value):
        return float(value) if value not in ("", None) else "—"

    labels = {("strong", "cv"): "強增強 / 交叉驗證", ("strong", "test"): "強增強 / 測試集(多數決)",
              ("none", "cv"): "無增強 / 交叉驗證", ("none", "test"): "無增強 / 測試集(多數決)"}
    header = ["前處理", "參數", "n", "mf Dice", "可辨識 mf Dice", "無法辨識 mf Dice", "命中 1R", "ian Dice", "差異(vs 原圖)", "95% CI 下界", "95% CI 上界", "Holm p", "判定"]
    for (setting, dataset), title in labels.items():
        rows = [r for r in data if r["setting"] == setting and r["dataset"] == dataset]
        if not rows:
            continue
        values = [[r["variant"], r["parameters"], int(r["n"]), num(r["mf_dice"]), num(r.get("mf_dice_visible")), num(r.get("mf_dice_not_visible")),
                   num(r["mf_hit_1R"]), num(r["ian_dice"]), num(r["diff"]), num(r["ci_low"]), num(r["ci_high"]), num(r["holm_p"]), r["verdict"] or "基準"]
                  for r in rows]
        s.table(f"{title}(Friedman p = {float(rows[0]['friedman_p']):.3g})", header, lambda first, values=values: values,
                [None, None, None, D4, D4, D4, PCT, D4, D4, D4, D4, P, None])
    s.cell(1, "來源: kfold_preprocessing.py", italic=True)
    s.finish()


PANO = ROOT / "compare_results" / "kfold_pano_mf_559"
PANO_SETTINGS = [("cv", "CV(559)"), ("test single", "測試 單一模型平均"), ("test ensemble", "測試 五折多數決")]
PANO_METRICS = ["mf_dice", "mf_hit_1R", "ian_dice", "mf_predicted"]


def read_summary(name):
    # kfold_pano 的 summary*.csv → {(model, setting): row}
    path = PANO / name
    if not path.exists():
        print(f"[略過] 找不到 {path}")
        return {}
    with open(path, newline="") as f:
        return {(r["model"], r["setting"]): r for r in csv.DictReader(f)}


def read_postproc():
    # mf_postproc.py 的 summary_postproc.csv,只取 lcc → {(model, setting): row}
    path = PANO / "summary_postproc.csv"
    if not path.exists():
        print(f"[略過] 找不到 {path}")
        return {}
    with open(path, newline="") as f:
        return {(r["model"], r["setting"]): r for r in csv.DictReader(f) if r["post"] == "lcc"}


def seed_rows(data, tag, seed_ids, metric, settings=PANO_SETTINGS):
    # 每個設定一列: 各種子的值 + 平均 / 標準差公式。種子 0 的模型名稱沒有 _s0
    cols = [get_column_letter(2 + i) for i in range(len(seed_ids))]
    out = []
    for setting, label in settings:
        vals = [data.get((tag + (f"_s{k}" if k else ""), setting), {}).get(metric) for k in seed_ids]
        out.append([label] + [float(v) if v else None for v in vals]
                   + [f"=AVERAGE({cols[0]}{{r}}:{cols[-1]}{{r}})", f"=STDEV({cols[0]}{{r}}:{cols[-1]}{{r}})"])
    return out


def bootstrap_table(s, name, title, note):
    # seed_bootstrap.py 的合併種子結果(seeds = pooled)
    path = PANO / name
    if not path.exists():
        print(f"[略過] 找不到 {path}")
        return
    with open(path, newline="") as f:
        boot = [r for r in csv.DictReader(f) if r["seeds"] == "pooled"]
    mlabels = {"mf_hit_1R": "mf 命中 1R", "mf_dice": "mf Dice", "ian_dice": "ian Dice"}
    s.table(title, ["集合", "指標", "ROI 數", "組數", "差", "95% CI 下", "95% CI 上", "p"],
            lambda first: [[r["setting"], mlabels[r["metric"]], int(r["n_roi"]), int(r["n_clusters"]), float(r["diff"]),
                            float(r["ci_low"]), float(r["ci_high"]), float(r["p"])] for r in boot],
            [None, None, None, None, D4, D4, D4, D3], note=note)


def read_rows(name):
    path = PANO / name
    if not path.exists():
        print(f"[略過] 找不到 {path}")
        return []
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def pano_more(wb):
    # 2026-10-30: 10/28–29 在新切分上的延伸實驗(都是 lcc 後)
    s = Sheet(wb, "新切分_延伸", "新切分的延伸實驗(頦孔一律 lcc)",
              ["困難案例、可辨識性、串接式推論、解析度 384、推論時翻轉、半監督;數字讀各程式的輸出 CSV,平均與標準差用公式"])
    unc = read_rows("uncertainty_summary_lcc.csv")
    s.table("以模型分歧標出困難案例(U-Net 五個種子)",
            ["集合", "ROI", "沒命中", "AUROC u→沒命中", "AUROC u→無法辨識", "命中 @100%", "@95%", "@90%", "@80%"],
            lambda first: [[r["setting"], int(r["n"]), int(r["misses"]), float(r["auroc_miss"]), float(r["auroc_not_visible"]),
                            float(r["hit@1.00"]), float(r["hit@0.95"]), float(r["hit@0.90"]), float(r["hit@0.80"])] for r in unc],
            [None, None, None, D3, D3, PCT, PCT, PCT, PCT],
            note="@x% = 只保留分歧度最低的 x%,其餘交給醫師。來源: pano_uncertainty.py --post lcc → uncertainty_summary_lcc.csv")

    vis = read_rows("visibility_strata_lcc.csv")
    s.table("可辨識性分層", ["模型", "集合", "分層", "ROI", "mf Dice", "Dice 95% CI 下", "Dice 95% CI 上", "mf 命中 1R", "命中 95% CI 下", "命中 95% CI 上"],
            lambda first: [[r["model"], r["setting"], r["stratum"], int(r["n"])] + [float(r[k]) for k in
                           ("mf_dice", "mf_dice_low", "mf_dice_high", "mf_hit_1R", "mf_hit_low", "mf_hit_high")] for r in vis],
            [None, None, None, None, D4, D4, D4, D4, D4, D4],
            note="來源: pano_visibility.py --post lcc → visibility_strata_lcc.csv;可辨識 = 嚴格標準下人眼看得出頦孔")

    cascade = read_rows("cascade_s01234_lcc.csv")
    policies = [("single", "", "單一模型"), ("cascade2", "0.5", "先 2 折,tau 0.5"), ("cascade3", "0.3", "先 3 折,tau 0.3"), ("full", "", "五折多數決")]
    for metric, mlabel in [("hit", "mf 命中 1R"), ("folds_run", "平均跑幾折")]:
        def rows(first, metric=metric):
            out = []
            for tag, name in [("r34_imagenet_aug", "U-Net"), ("yolo26s-seg", "YOLO26s")]:
                for policy, tau, plabel in policies:
                    vals = {int(r["seed"]): float(r[metric]) for r in cascade
                            if r["model"] == tag and r["policy"] == policy and r["tau"] == tau}
                    out.append([f"{name} {plabel}"] + [vals.get(k) for k in range(5)] + ["=AVERAGE(B{r}:F{r})", "=STDEV(B{r}:F{r})"])
            return out
        s.table(f"串接式推論(測試集): {mlabel}", ["策略", "種子 0", "種子 1", "種子 2", "種子 3", "種子 4", "平均", "標準差"], rows,
                [None] + [D4 if metric == "hit" else "0.00"] * 7,
                note="lcc 只套在最後輸出;投票與一致性判斷用原遮罩。來源: cascade_infer.py --post lcc → cascade_s01234_lcc.csv")

    post = read_postproc()
    for tag, name in [("r34_imagenet_aug", "R34/512(主結果)"), ("r34_imagenet_aug_tta", "R34/512 + 翻轉"),
                      ("r34_imagenet_aug_384", "R34/384"), ("r34_imagenet_aug_384_tta", "R34/384 + 翻轉")]:
        # _tta 的種子名稱是 <設定>_s<k>_tta
        base, tta = (tag[:-4], "_tta") if tag.endswith("_tta") else (tag, "")
        data = {(base + (f"_s{k}" if k else ""), setting): post.get((base + (f"_s{k}" if k else "") + tta, setting), {})
                for k in range(5) for setting, _ in PANO_SETTINGS}
        s.table(f"{name}: mf 命中 1R,lcc(五個訓練種子)", ["設定", "種子 0", "種子 1", "種子 2", "種子 3", "種子 4", "平均", "標準差"],
                lambda first, base=base, data=data: seed_rows(data, base, range(5), "mf_hit_1R"), [None] + [D4] * 7,
                note="翻轉 = 推論時原圖 + 左右翻轉,softmax 平均(2 倍運算),不採用。來源: unet_tta.py、mf_postproc.py → summary_postproc.csv")
    bootstrap_table(s, "seed_bootstrap_r34_imagenet_aug_384_s01234_lcc.csv", "R34/512 − R34/384(合併五個種子,lcc)",
                    "來源: seed_bootstrap.py --post lcc --prefix _r384 --other r34_imagenet_aug_384 --seeds 0 1 2 3 4")

    for tag, name in [("r34_imagenet_aug", "基準(只用真標註)"), ("r34_imagenet_aug_pl", "半監督(加假標註)")]:
        s.table(f"{name}: mf 命中 1R,lcc(種子 0–2)", ["設定", "種子 0", "種子 1", "種子 2", "平均", "標準差"],
                lambda first, tag=tag: seed_rows(post, tag, range(3), "mf_hit_1R"), [None] + [D4] * 5,
                note="假標註: 沒有標註的 1018 個 ROI 中,多數決有頦孔且 u ≤ 0.2577 的約 26%。來源: kfold_pseudo.py、mf_postproc.py")
    bootstrap_table(s, "seed_bootstrap_r34_imagenet_aug_pl_s012_lcc.csv", "基準 − 半監督(合併三個種子,lcc)",
                    "來源: seed_bootstrap.py --post lcc --prefix _pl --other r34_imagenet_aug_pl --seeds 0 1 2")

    # 2026-10-30: 定數字前的檢查
    ci = read_rows("uncertainty_ci_lcc.csv")
    s.table("困難案例的 95% CI(以組 bootstrap)", ["集合", "指標", "估計", "CI 下", "CI 上"],
            lambda first: [[r["setting"], r["metric"], float(r["estimate"]), float(r["ci_low"]), float(r["ci_high"])] for r in ci],
            [None, None, D4, D4, D4], note="gain = 保留後命中 − 全部命中。來源: uncertainty_ci.py → uncertainty_ci_lcc.csv")
    threshold = read_rows("hit_threshold_lcc.csv")
    s.table("命中門檻的敏感度(U-Net vs YOLO26s,五個種子)", ["集合", "門檻 (R)", "ROI", "U-Net", "YOLO26s", "差", "CI 下", "CI 上", "p"],
            lambda first: [[r["setting"], float(r["threshold_R"]), int(r["n"])] + [float(r[k]) for k in ("unet", "yolo26s", "diff", "ci_low", "ci_high", "p")] for r in threshold],
            [None, "0.0", None, D4, D4, D4, D4, D4, D4], note="來源: hit_threshold.py → hit_threshold_lcc.csv")
    single = read_rows("distill/single_r18_imagenet_aug_384_lcc_summary.csv")
    s.table("R18/384 單一模型對五折多數決(測試集)", ["項目", "命中或差", "標準差", "mf Dice", "CI 下", "CI 上", "p"],
            lambda first: [[r["row"]] + [float(r[k]) if r[k] else None for k in ("hit", "hit_sd", "mf_dice", "ci_low", "ci_high", "p")] for r in single],
            [None, D4, D4, D4, D4, D4, D4], note="單一模型用全部 559 張訓練,種子 0–4;取代舊的「90.3%、少 2.3 點」。來源: single_small_lcc.py")
    s.finish()


def pano(wb):
    # 2026-10-28: 新切分(全景片層級,kfold_pano.py)的結果。之前的工作表都是舊的病人層級切分(kfold_patient)
    s = Sheet(wb, "新切分_主結果", "新切分(病人 ∪ 來源全景片,301 組)的結果",
              ["CV 559 張(每張用沒看過它的那一折);測試 137 張(去掉與訓練同一張全景片的 209_Left_roi)",
               "論文主表以這張工作表為準;其他工作表是舊的病人層級切分"])
    main = read_summary("summary.csv")
    b8 = read_summary("summary_b8.csv")
    names = [("r34_imagenet_aug", "U-Net ResNet34"), ("yolov8n-seg", "YOLOv8n-seg"), ("yolov8s-seg", "YOLOv8s-seg"),
             ("yolo11n-seg", "YOLO11n-seg"), ("yolo11s-seg", "YOLO11s-seg"), ("yolo26n-seg", "YOLO26n-seg"),
             ("yolo26s-seg", "YOLO26s-seg(batch 16)")]
    header = ["模型", "設定", "mf Dice", "mf 命中 1R", "ian Dice", "mf 有預測"]
    def rows(first):
        out = []
        for setting, label in PANO_SETTINGS:
            for tag, name in names:
                r = main.get((tag, setting))
                if r:
                    out.append([name, label] + [float(r[m]) for m in PANO_METRICS])
            r = b8.get(("yolo26s-seg_b8", setting))
            if r:
                out.append(["YOLO26s-seg(batch 8,補充)", label] + [float(r[m]) for m in PANO_METRICS])
        return out
    s.table("七個模型(種子 0)", header, rows, [None, None, D4, D4, D4, D4],
            note="來源: kfold_pano.py → summary.csv、summary_b8.csv。YOLO 推論每類取信心最高的候選")

    # 五個訓練種子: U-Net 與 YOLO26s,平均與標準差用公式
    seeds = [("", 0)] + [(f"_s{k}", k) for k in range(1, 5)]
    data = {sfx: read_summary(f"summary{sfx}.csv") for sfx, _ in seeds}
    for tag, name in [("r34_imagenet_aug", "U-Net ResNet34"), ("yolo26s-seg", "YOLO26s-seg(batch 16)")]:
        for metric, mlabel in [("mf_hit_1R", "mf 命中 1R"), ("mf_dice", "mf Dice"), ("ian_dice", "ian Dice")]:
            def rows(first, tag=tag, metric=metric):
                out = []
                for setting, label in PANO_SETTINGS:
                    vals = [data[sfx].get((tag + sfx, setting), {}).get(metric) for sfx, _ in seeds]
                    vals = [float(v) if v else None for v in vals]
                    out.append([label] + vals + ["=AVERAGE(B{r}:F{r})", "=STDEV(B{r}:F{r})"])
                return out
            s.table(f"{name}: {mlabel}(五個訓練種子)", ["設定", "種子 0", "種子 1", "種子 2", "種子 3", "種子 4", "平均", "標準差"],
                    rows, [None] + [D4] * 7)

    bootstrap_table(s, "seed_bootstrap_yolo26s-seg_s01234.csv", "U-Net − YOLO26s(合併五個種子,以病人 / 組為單位,無後處理)",
                    "組 bootstrap 20000 次;p 為組為單位的符號翻轉置換檢定。來源: seed_bootstrap.py --seeds 0 1 2 3 4")

    # 2026-10-28 起頦孔只留最大連通區塊(lcc)列為標準流程,論文措辭以下面兩張表為準
    post = read_postproc()
    for tag, name in [("r34_imagenet_aug", "U-Net ResNet34"), ("yolo26s-seg", "YOLO26s-seg(batch 16)")]:
        for metric, mlabel in [("mf_hit_1R", "mf 命中 1R"), ("mf_dice", "mf Dice")]:
            s.table(f"{name}: {mlabel},lcc(五個訓練種子)", ["設定", "種子 0", "種子 1", "種子 2", "種子 3", "種子 4", "平均", "標準差"],
                    lambda first, tag=tag, metric=metric: seed_rows(post, tag, range(5), metric), [None] + [D4] * 7,
                    note="來源: mf_postproc.py → summary_postproc.csv(post = lcc)。YOLO 推論本來就只取一個實例,lcc 不改變結果")
    bootstrap_table(s, "seed_bootstrap_yolo26s-seg_s01234_lcc.csv", "U-Net − YOLO26s(合併五個種子,以病人 / 組為單位,lcc;論文主結果)",
                    "來源: seed_bootstrap.py --post lcc --seeds 0 1 2 3 4")
    s.finish()

    # 精度對運算量: 單一 U-Net 的 encoder 與輸入尺寸,三個種子
    s = Sheet(wb, "新切分_運算量", "精度對運算量(U-Net,新切分)",
              ["GMAC = 每個 Conv2d 的乘加數加總(batch 1);每個設定三個訓練種子(0–2),平均與標準差用公式計算"])
    configs = [("r34_imagenet_aug", "ResNet34 / 512", 40.91), ("r18_imagenet_aug", "ResNet18 / 512", 31.24),
               ("r34_imagenet_aug_384", "ResNet34 / 384", 23.01), ("r18_imagenet_aug_384", "ResNet18 / 384", 17.57)]
    for setting, label in PANO_SETTINGS:
        def rows(first, setting=setting):
            out = []
            for tag, name, gmac in configs:
                # 主結果 r34_imagenet_aug 的種子在 summary_s<k>.csv;其他設定在 summary_<設定>[_s<k>].csv
                vals = []
                for k in range(3):
                    sfx = f"_s{k}" if k else ""
                    file = f"summary{sfx}.csv" if tag == "r34_imagenet_aug" else f"summary_{tag}{sfx}.csv"
                    r = (data.get(sfx) if tag == "r34_imagenet_aug" else read_summary(file)).get((tag + sfx, setting))
                    vals.append(float(r["mf_hit_1R"]) if r else None)
                n = sum(v is not None for v in vals)
                out.append([name, gmac] + vals + ["=AVERAGE(C{r}:E{r})", "=STDEV(C{r}:E{r})" if n >= 2 else "—"])
            return out
        s.table(f"mf 命中 1R: {label}", ["設定", "GMAC", "種子 0", "種子 1", "種子 2", "平均", "標準差"], rows,
                [None, "0.00"] + [D4] * 5,
                note="來源: kfold_pano.py --models unet --unet-config <設定> [--seed k];ResNet18 / 512 只有種子 0")
    s.finish()


def others(wb):
    s = Sheet(wb, "其他實驗", "其他已測試的方向")
    s.table("YOLO 推論規則(影像層級切分)", ["規則", "CV mf Dice", "CV ian Dice", "測試 mf Dice", "測試 ian Dice"], lambda first: [
        ("固定門檻 0.4", 0.5952, 0.5840, 0.4888, 0.6013),
        ("每類取信心最高", 0.6720, 0.6188, 0.6320, 0.6280),
        ("取最高 + U-Net 補漏", 0.6805, 0.6190, 0.6490, 0.6281),
    ], [None, D4, D4, D4, D4], note="已採用「每類取信心最高」。來源: eval_kfold_variants.py")
    s.table("候選重新挑選(YOLO)", ["方法", "CV mf Dice", "CV ian Dice", "測試 mf Dice", "測試 ian Dice"], lambda first: [
        ("每類取信心最高", 0.6720, 0.6188, 0.6320, 0.6280),
        ("上限:前 50 名中最佳(偷看答案)", 0.7092, 0.6401, 0.6573, 0.6429),
        ("mf–ian 距離一起挑(tau 40)", 0.6824, 0.6161, 0.6414, 0.6272),
    ], [None, D4, D4, D4, D4], note="上限只有 +0.02–0.04,不採用。來源: rerank_candidates.py")
    s.table("YOLO 模型大小", ["模型", "CV mf Dice", "CV ian Dice", "測試 mf Dice", "測試 ian Dice"], lambda first: [
        ("yolov8n-seg", 0.6720, 0.6188, 0.6320, 0.6280),
        ("yolov8s-seg", 0.6754, 0.6091, 0.6422, 0.6223),
    ], [None, D4, D4, D4, D4], note="無顯著差異。來源: kfold_yolo_model.py")
    s.table("U-Net 設計變體(病人層級)", ["版本", "CV mf Dice", "CV ian Dice", "測試 mf Dice", "測試 ian Dice"], lambda first: [
        ("最終模型", 0.7005, 0.6427, 0.6697, 0.6517),
        ("右側鏡射成左側(canonical)", 0.7034, 0.6362, 0.6922, 0.6690),
        ("座標通道(coord)", 0.6970, 0.6462, 0.6801, 0.6545),
        ("兩者一起", 0.6912, 0.6465, 0.6794, 0.6571),
        ("只訓練頦孔(mf_only)", 0.7111, "—", 0.6882, "—"),
    ], [None, D4, D4, D4, D4], note="皆無顯著差異。來源: kfold_patient_variants.py")
    s.table("切分方式的影響(交叉驗證)", ["模型", "影像層級 mf", "病人層級 mf", "影像層級 ian", "病人層級 ian", "mf 差異"], lambda first: [
        ("U-Net ResNet34", 0.7054, 0.7005, 0.6483, 0.6427, "=C{r}-B{r}"),
        ("YOLOv8n-seg", 0.6720, 0.6783, 0.6188, 0.6111, "=C{r}-B{r}"),
    ], [None, D4, D4, D4, D4, D4], note="左右側跨折洩漏沒有造成可測量的虛高。來源: kfold_unet_variants.py、kfold_patient.py")
    s.table("前處理比較(YOLO,397 張,fold 0–3 平均)", ["前處理", "mask mAP50-95"], lambda first: [
        ("original", 0.2535), ("gamma", 0.2575), ("bilateral_filter", 0.2556), ("bf_gamma", 0.2504),
        ("gamma_clahe", 0.2555), ("bf_gamma_clahe", 0.2520), ("gamma_clahe_unsharp_mask", 0.2564), ("bf_gamma_clahe_um", 0.2530),
    ], [None, D4], note="差異遠小於折間波動(0.22–0.28)。由 compare_mf_397_5fold/runs.csv 的 fold 0–3 計算")
    s.finish()


def evidence(wb):
    s = Sheet(wb, "可信度", "各實驗的可信度",
              ["A:病人層級切分、最終設定或嚴格控制的比較、有獨立測試集,可直接寫入論文",
               "B:有一項明確的限制,可作輔助證據並說明限制;C:條件不同、未完成或已被取代,不建議引用數字"])
    rows = [
        ("U-Net 與 YOLO 比較", "A", "病人層級", "最終設定", "有", 3, "主要結果", "seed_replicates.py"),
        ("完整指標與信賴區間", "A", "病人層級", "最終設定", "有", 1, "以病人為單位重抽樣", "paper_metrics.py"),
        ("消融實驗", "A", "病人層級", "最終設定與變體", "有", 1, "設計間差距遠大於種子波動", "ablation_patient.py"),
        ("學習曲線", "A", "病人層級", "最終設定", "有", 1, "更新步數相同", "learning_curve.py"),
        ("兩批資料的表現差異", "A", "病人層級", "6 種模型", "—", 1, "描述性結果,所有模型一致", "paper_metrics.py、batch_analysis.py"),
        ("U-Net 設計變體", "A", "病人層級", "最終設定與變體", "有", 1, "陰性結果", "kfold_patient_variants.py"),
        ("測試集獨立性檢查", "A", "—", "—", "—", "—", "以影像內容找出每張 ROI 的來源全景片(614/697 可驗證):測試集與訓練資料無共用全景片,訓練資料內無重複裁切或跨折", "locate_roi_sources.py"),
        ("前處理比較(新)", "A", "病人層級", "最終設定", "有", 1, "參數與統計方法事先固定;強增強與無增強兩種設定", "kfold_preprocessing.py"),
        ("交叉批次實驗", "B", "病人層級", "最終設定", "有", 1, "單一批次只用約一半資料", "cross_batch.py"),
        ("頦孔可辨識性分層", "B", "病人層級", "最終設定與 YOLO", "有", 1, "單一判定者,與先前判斷的 kappa 0.56", "visibility_grading.py"),
        ("人工標註一致性", "B", "—", "—", "—", "—", "只有 40 張,單一標註者前後兩次", "annotation_agreement.py"),
        ("頦孔暗度分析", "B", "病人層級", "最終設定", "—", 1, "暗度指標是近似值,兩種算法結論不同", "check_labels.py、contrast_analysis.py"),
        ("YOLO 推論規則、候選挑選", "B", "影像層級", "YOLOv8n", "有", 1, "切分方式已證實影響在 0.01 以內", "eval_kfold_variants.py、rerank_candidates.py"),
        ("YOLOv8s 與 v8n", "B", "影像層級", "YOLO", "有", 1, "同上", "kfold_yolo_model.py"),
        ("五種 U-Net 設定(影像層級)", "C", "影像層級", "多種", "有", 1, "已被病人層級的消融實驗取代", "kfold_unet_variants.py"),
        ("前處理比較(舊)", "C", "影像層級", "YOLO", "無", 1, "397 張、mAP 指標、fold 4 未完成;由新實驗取代", "compare_preprocessing.py"),
        ("從零訓練的 U-Net 與 YOLO 比較", "C", "影像層級", "原 U-Net", "有", 1, "U-Net 未調整,結論已被推翻", "kfold_yolo_unet.py"),
        ("舊 U-Net 權重的測試結果", "C", "不明", "原 U-Net", "有", "—", "訓練過程與資料無法查證", "predict_unet.py"),
    ]
    s.table("實驗清單", ["實驗", "等級", "切分", "模型", "測試集", "種子數", "說明", "程式"], lambda first: rows)
    s.finish()


def appendix(wb):
    files = [
        ("附錄_paper_metrics", ROOT / "compare_results" / "kfold_patient_mf_559" / "paper_metrics.csv"),
        ("附錄_learning_curve", ROOT / "compare_results" / "learning_curve_mf_559" / "learning_curve.csv"),
        ("附錄_agreement", ROOT / "annotation_agreement" / "agreement.csv"),
    ]
    for title, path in files:
        if not path.exists():
            print(f"[略過] 找不到 {path}")
            continue
        ws = wb.create_sheet(title)
        with open(path, newline="") as f:
            for i, row in enumerate(csv.reader(f), start=1):
                for j, value in enumerate(row, start=1):
                    try:
                        value = float(value)
                    except ValueError:
                        pass
                    c = ws.cell(row=i, column=j, value=value)
                    c.font = Font(name=FONT, bold=i == 1, size=10)
        ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for build in (overview, pano, pano_more, seeds,diff_ci, details, ablation, learning, batches, agreement, visibility, preprocessing, others, evidence, appendix):
        build(wb)
    # 開啟時重新計算所有公式
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"已輸出: {OUT}")


if __name__ == "__main__":
    main()
